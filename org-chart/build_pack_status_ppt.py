#!/usr/bin/env python3
"""Build the Design-to-Tender pack-status PowerPoint (workflow + summary tracker)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

DIR = "/home/user/phd-tracker/org-chart"
NAVY = RGBColor(0x0B,0x2E,0x63); BLUE = RGBColor(0x1F,0x5F,0xD6)
INK = RGBColor(0x0F,0x1B,0x3D); MUTED = RGBColor(0x5B,0x6B,0x8C); WHITE = RGBColor(0xFF,0xFF,0xFF)
# stage -> (name, light fill)
STAGE = {
 1:("Pre-IDC Pack",              RGBColor(0xEC,0xF3,0xFF)),
 2:("IDC",                        RGBColor(0xE1,0xEC,0xFB)),
 3:("Full ITT Pack — Rev 01",     RGBColor(0xD6,0xE4,0xF7)),
 4:("TII Instruction to Issue C01",RGBColor(0xE7,0xE7,0xFA)),
 5:("AWDS C01 Submission",        RGBColor(0xDE,0xDE,0xF6)),
 6:("Ready for Tender",           RGBColor(0xDD,0xE8,0xF6)),
 7:("ITT Issue Date",             RGBColor(0xD2,0xE0,0xF7)),
}
# 16 contracts with indicative current stage (confirm against live programme)
CONTRACTS = [
 ("M140", 2), ("M141", 2), ("M142", 1), ("M145", 1),
 ("M146", 1), ("M147", 1), ("M130", 1), ("M135", 1),
 ("M138", 1), ("M160", 3), ("M161", 1), ("M190", 2),
 ("M111", 1), ("Minor Works", 1), ("Env. Baseline", 1), ("DAA Works", 1),
]

prs = Presentation()
prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

def img_wh(path):
    from PIL import Image
    w,h = Image.open(path).size; return w/h

# ---------------- Slide 1 : workflow image ----------------
s1 = prs.slides.add_slide(blank)
s1.shapes.add_picture(f"{DIR}/programme-flow.png", 0, 0, width=prs.slide_width, height=prs.slide_height)

# ---------------- Slide 2 : pack status summary ----------------
s2 = prs.slides.add_slide(blank)

# MetroLink logo top-left
ar_ml = img_wh(f"{DIR}/logos/metrolink.png")
s2.shapes.add_picture(f"{DIR}/logos/metrolink.png", Inches(0.5), Inches(0.34), height=Inches(0.42))

def textbox(slide, l,t,w,h):
    tb = slide.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h)); return tb.text_frame

# title
tf = textbox(s2, 3.4, 0.30, 9.4, 0.9)
p = tf.paragraphs[0]; p.alignment = PP_ALIGN.RIGHT
r = p.add_run(); r.text = "Design-to-Tender Workflow — Pack Status Summary"
r.font.size = Pt(22); r.font.bold = True; r.font.color.rgb = NAVY
p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.RIGHT
r2 = p2.add_run(); r2.text = "Current workflow stage of each contract package"
r2.font.size = Pt(12); r2.font.italic = True; r2.font.color.rgb = MUTED

# legend line
lf = textbox(s2, 0.5, 1.18, 12.33, 0.5)
lp = lf.paragraphs[0]
lead = lp.add_run(); lead.text = "Workflow stages:   "
lead.font.size = Pt(10.5); lead.font.bold = True; lead.font.color.rgb = NAVY
parts = " · ".join([f"{n} {STAGE[n][0]}" for n in range(1,8)])
lr = lp.add_run(); lr.text = parts
lr.font.size = Pt(10.5); lr.font.color.rgb = INK

# two tables of 8
def make_table(rows_data, left):
    nrows = len(rows_data)+1
    tbl_shape = s2.shapes.add_table(nrows, 2, Inches(left), Inches(1.75),
                                    Inches(6.0), Inches(0.42*nrows))
    tbl = tbl_shape.table
    tbl.columns[0].width = Inches(1.7); tbl.columns[1].width = Inches(4.3)
    # header
    for j,txt in enumerate(["Contract","Design-to-Tender Workflow Status"]):
        c = tbl.cell(0,j); c.fill.solid(); c.fill.fore_color.rgb = NAVY
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        pr = c.text_frame.paragraphs[0]; run = pr.add_run(); run.text = txt
        run.font.bold = True; run.font.size = Pt(11); run.font.color.rgb = WHITE
        c.margin_left = Inches(0.08); c.margin_top = Inches(0.02); c.margin_bottom = Inches(0.02)
    for i,(code,stage) in enumerate(rows_data, start=1):
        name, fill = STAGE[stage]
        c0 = tbl.cell(i,0); c0.fill.solid(); c0.fill.fore_color.rgb = WHITE
        c0.vertical_anchor = MSO_ANCHOR.MIDDLE
        p0 = c0.text_frame.paragraphs[0]; r0 = p0.add_run(); r0.text = code
        r0.font.bold = True; r0.font.size = Pt(11); r0.font.color.rgb = NAVY
        c1 = tbl.cell(i,1); c1.fill.solid(); c1.fill.fore_color.rgb = fill
        c1.vertical_anchor = MSO_ANCHOR.MIDDLE
        p1 = c1.text_frame.paragraphs[0]
        rn = p1.add_run(); rn.text = f"{stage}  "; rn.font.bold = True; rn.font.size = Pt(11); rn.font.color.rgb = BLUE
        rn2 = p1.add_run(); rn2.text = f"·  {name}"; rn2.font.size = Pt(11); rn2.font.color.rgb = INK
        for c in (c0,c1):
            c.margin_left = Inches(0.08); c.margin_top = Inches(0.02); c.margin_bottom = Inches(0.02)
    # row heights
    for r_ in tbl.rows: r_.height = Inches(0.4)
    return tbl

make_table(CONTRACTS[:8], 0.5)
make_table(CONTRACTS[8:], 6.83)

# note
nf = textbox(s2, 0.5, 6.55, 9.5, 0.5)
np_ = nf.paragraphs[0]
nr = np_.add_run()
nr.text = "Status indicative as at 30-Jun-2026 — to be confirmed and maintained against the live P6 integrated programme. Stage numbers correspond to the Design-to-Tender Workflow."
nr.font.size = Pt(9.5); nr.font.italic = True; nr.font.color.rgb = MUTED

# footer logos
s2.shapes.add_picture(f"{DIR}/logos/atkinsrealis.png", Inches(10.0), Inches(6.95), height=Inches(0.3))
s2.shapes.add_picture(f"{DIR}/logos/tetratech.png", Inches(12.1), Inches(6.9), height=Inches(0.36))

out = f"{DIR}/AWDS_Pack_Status_Summary.pptx"
prs.save(out)
print("saved", out)
